#!/usr/bin/env python3
"""
Build "The Treaty That Failed Twice" — 7 slides, 16:9, 5 minutes.

Design system is defined once at the top and every slide obeys it.
Run:  python3 build.py
Out:  The-Treaty-That-Failed-Twice.pptx
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

OUT = "The-Treaty-That-Failed-Twice.pptx"

# Slide 5 reveals one row per click. It is the only animation in the deck,
# which is what makes it read as deliberate rather than decorative.
# Set False to ship a static slide 5 instead.
ANIMATE = True

# ---------------------------------------------------------------- palette
# One hue family for the accent. Amber on navy and amber on bone are the same
# hue (15 deg); they differ in lightness only, because a single value cannot
# clear 4.5:1 against both backgrounds at once.
NAVY        = RGBColor(0x0B, 0x1B, 0x2B)   # dominant, 65%
BONE        = RGBColor(0xF2, 0xEF, 0xE9)   # support, 25%
AMBER       = RGBColor(0xE8, 0x66, 0x3A)   # accent on dark   5.33:1 on navy
AMBER_DEEP  = RGBColor(0xB0, 0x3E, 0x17)   # accent on light  5.19:1 on bone
SLATE       = RGBColor(0x4A, 0x6B, 0x7C)   # muted text on light  5.02:1
SLATE_LIGHT = RGBColor(0x8F, 0xA9, 0xB8)   # muted text on dark   7.05:1
RULE        = RGBColor(0x6E, 0x87, 0x97)   # hairlines on bone    3.29:1

# ---------------------------------------------------------------- typography
SERIF = "Cambria"    # headings
SANS  = "Calibri"    # body
# roles -> size (pt)
TITLE, BIGSTAT, ROW, QUOTE, HEAD, BODY, LABEL, CAPTION = 40, 88, 20, 24, 28, 16, 13, 10
LINE = 1.3

# ---------------------------------------------------------------- grid
M      = 0.6          # safe margin, every edge
COLW   = 0.827778
GUT    = 0.2
PITCH  = COLW + GUT
def col(n):            return M + (n - 1) * PITCH
def span(a, b):        return (b - a + 1) * COLW + (b - a) * GUT
FULL   = span(1, 12)   # 12.1333
BOTTOM = 7.5 - M       # 6.90 — nothing crosses this

prs = Presentation()
prs.slide_width  = Emu(12192000)   # exactly 13.333in; Inches(13.3333) is 31 EMU short
prs.slide_height = Emu(6858000)    # exactly 7.5in
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def strip_style(shape):
    """Autoshapes carry a <p:style> effectRef that re-introduces the theme
    shadow even when effectLst is empty. Fill and line are set explicitly
    everywhere, so the style reference is pure liability."""
    el = shape._element.find(qn("p:style"))
    if el is not None:
        shape._element.remove(el)


def alt(shape, text):
    """Alt text for screen readers."""
    try:
        shape._element._nvXxPr.cNvPr.set("descr", text)
    except AttributeError:
        pass


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def text(slide, x, y, w, h, runs, size=BODY, font=SANS, color=NAVY, bold=False,
         italic=False, align=PP_ALIGN.LEFT, line=LINE, caps=False, spc=None,
         anchor=MSO_ANCHOR.TOP, space_after=0, descr=None):
    """runs: str, or list of (text, {overrides}) tuples, or list of such lists
    (one inner list per paragraph)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    if isinstance(runs, str):
        paras = [[(runs, {})]]
    elif runs and isinstance(runs[0], tuple):
        paras = [runs]
    else:
        paras = runs

    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        if space_after:
            p.space_after = Pt(space_after)
        for txt, ov in para:
            r = p.add_run()
            r.text = txt.upper() if ov.get("caps", caps) else txt
            f = r.font
            f.name = ov.get("font", font)
            f.size = Pt(ov.get("size", size))
            f.bold = ov.get("bold", bold)
            f.italic = ov.get("italic", italic)
            f.color.rgb = ov.get("color", color)
            s = ov.get("spc", spc)
            if s:
                r._r.get_or_add_rPr().set("spc", str(int(s)))
            if ov.get("strike"):
                r._r.get_or_add_rPr().set("strike", "sngStrike")
    if descr:
        alt(box, descr)
    return box


def rect(slide, x, y, w, h, fill=None, line_color=None, line_w=1.0,
         shape=MSO_SHAPE.RECTANGLE, descr=None, adj=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line_color is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line_color
        s.line.width = Pt(line_w)
    if adj:
        for i, v in enumerate(adj):
            s.adjustments[i] = v
    s.shadow.inherit = False
    strip_style(s)
    if s.has_text_frame:
        s.text_frame.text = ""
    alt(s, descr or "")
    return s


def line_h(slide, x1, x2, y, color, w=1.0):
    c = slide.shapes.add_connector(1, Inches(x1), Inches(y), Inches(x2), Inches(y))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    strip_style(c)
    return c


def line_v(slide, x, y1, y2, color, w=1.0):
    c = slide.shapes.add_connector(1, Inches(x), Inches(y1), Inches(x), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    strip_style(c)
    return c


def dot(slide, cx, cy, d, color):
    return rect(slide, cx - d / 2, cy - d / 2, d, d, fill=color,
                shape=MSO_SHAPE.OVAL)


def eyebrow(slide, x, y, w, txt, color, align=PP_ALIGN.LEFT):
    return text(slide, x, y, w, 0.26, txt, size=LABEL, font=SANS, color=color,
                bold=True, caps=True, spc=LABEL * 8, align=align)  # +8% track


def add_click_build(slide, groups):
    """Appear-on-click entrance for each group of shapes. python-pptx has no
    animation API, so the <p:timing> tree is written by hand. One click per
    group; every shape inside a group appears together."""
    from lxml import etree
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"

    def el(parent, tag, **attrs):
        t = "{%s}%s" % (P, tag)
        if parent is None:
            return etree.Element(t, **attrs)
        return etree.SubElement(parent, t, **attrs)

    ids = iter(range(3, 500))

    timing = el(None, "timing")
    tnLst  = el(timing, "tnLst")
    root   = el(el(tnLst, "par"), "cTn", id="1", dur="indefinite",
                restart="never", nodeType="tmRoot")
    seq    = el(el(root, "childTnLst"), "seq", concurrent="1", nextAc="seek")
    mainTn = el(seq, "cTn", id="2", dur="indefinite", nodeType="mainSeq")
    mainLst = el(mainTn, "childTnLst")

    for group in groups:
        click = el(el(mainLst, "par"), "cTn", id=str(next(ids)), fill="hold")
        el(el(click, "stCondLst"), "cond", delay="indefinite")
        inner = el(el(el(click, "childTnLst"), "par"), "cTn",
                   id=str(next(ids)), fill="hold")
        el(el(inner, "stCondLst"), "cond", delay="0")
        innerLst = el(inner, "childTnLst")
        for i, shape in enumerate(group):
            eff = el(el(innerLst, "par"), "cTn", id=str(next(ids)),
                     presetID="1", presetClass="entr", presetSubtype="0",
                     fill="hold", grpId="0",
                     nodeType="clickEffect" if i == 0 else "withEffect")
            el(el(eff, "stCondLst"), "cond", delay="0")
            st = el(el(eff, "childTnLst"), "set")
            beh = el(st, "cBhvr")
            btn = el(beh, "cTn", id=str(next(ids)), dur="1", fill="hold")
            el(el(btn, "stCondLst"), "cond", delay="0")
            el(el(beh, "tgtEl"), "spTgt", spid=str(shape.shape_id))
            el(el(beh, "attrNameLst"), "attrName").text = "style.visibility"
            el(el(st, "to"), "strVal", val="visible")

    for evt, tag in (("onPrev", "prevCondLst"), ("onNext", "nextCondLst")):
        cond = el(el(seq, tag), "cond", evt=evt, delay="0")
        el(el(cond, "tgtEl"), "sldTgt")

    bld = el(timing, "bldLst")
    for group in groups:
        for shape in group:
            el(bld, "bldP", spid=str(shape.shape_id), grpId="0")

    slide._element.append(timing)


def source(slide, txt, x=M, y=6.60, w=FULL, color=SLATE, align=PP_ALIGN.LEFT):
    return text(slide, x, y, w, 0.28, txt, size=CAPTION, color=color,
                line=1.1, align=align)


# ================================================================ SLIDE 1
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
# Slate, not amber: the accent is spent on the projected bar, and a section
# label is branding rather than an idea.
eyebrow(s, M, 0.60, 6.0, "Global plastics treaty", SLATE_LIGHT)

# Two honest data points, no axis furniture, labels sitting on the bars.
BASE, MAXH = 4.30, 3.00
h19 = MAXH * 460 / 1231
rect(s, 8.60, BASE - h19, 1.10, h19, fill=SLATE_LIGHT,
     descr="Bar: 460 million tonnes of plastic produced in 2019")
rect(s, 10.30, BASE - MAXH, 1.10, MAXH, fill=AMBER,
     descr="Bar: 1,231 million tonnes of plastic projected for 2060")
text(s, 8.60, BASE - h19 - 0.52, 1.60, 0.40, "460 Mt", size=ROW, font=SERIF,
     bold=True, color=BONE)
text(s, 10.30, BASE - MAXH - 0.52, 2.30, 0.40, "1,231 Mt", size=ROW,
     font=SERIF, bold=True, color=AMBER)
text(s, 8.60, BASE + 0.10, 1.60, 0.26, "2019", size=CAPTION, color=SLATE_LIGHT)
text(s, 10.30, BASE + 0.10, 2.30, 0.26, "2060, projected", size=CAPTION,
     color=SLATE_LIGHT)
source(s, "Source: OECD Global Plastics Outlook, 2022", x=8.60, y=4.85, w=4.13,
       color=RULE)

text(s, M, 5.55, 9.00, 0.78, "The treaty that failed twice",
     size=TITLE, font=SERIF, bold=True, color=BONE, line=1.1)
# Subtitle and byline share a box height and a bottom anchor so their
# baselines land on the same line despite the 16pt/10pt size difference.
text(s, M, 6.30, 7.00, 0.40, "And what that did to a movement",
     size=BODY, color=SLATE_LIGHT, line=1.1, anchor=MSO_ANCHOR.BOTTOM)
text(s, 8.00, 6.30, 3.40, 0.40, "Cesco Cugliari", size=CAPTION,
     color=SLATE_LIGHT, align=PP_ALIGN.RIGHT, line=1.1,
     anchor=MSO_ANCHOR.BOTTOM)
notes(s, "Plastic production is set to nearly triple by 2060 - 460 million "
         "tonnes in 2019 to a projected 1,231. In 2022, 175 countries agreed "
         "to write a treaty to stop that. They have now failed twice. What I "
         "want to show you is not the failure. It is what the failure did to "
         "the people fighting it.")

# ================================================================ SLIDE 2
s = prs.slides.add_slide(BLANK)
bg(s, BONE)
text(s, M, 0.60, FULL, 0.72, "One resolution, five sessions, no treaty",
     size=TITLE, font=SERIF, bold=True, color=NAVY, line=1.1)
# Brackets are the motif; the accent colour is reserved for the slide's one
# idea, which here is the two collapses.
text(s, M, 1.85, span(1, 6), 1.10, [
        [("In March 2022, 175 countries agreed to negotiate a binding treaty "
          "covering the ", {}),
         ("[full life cycle]", {"bold": True}),
         (" of plastic. Production, use, disposal.", {})]],
     size=BODY, color=NAVY)
text(s, M, 2.85, span(1, 6), 0.60,
     "That deadline came and went. So did the next one.", size=BODY,
     color=NAVY)

TL = [("Mar 2022",  "UNEA Resolution 5/14 adopted", False),
      ("2022-2024", "INC-1 to INC-4, four rounds",  False),
      ("Dec 2024",  "Busan, INC-5, collapsed",      True),
      ("Aug 2025",  "Geneva, INC-5.2, collapsed",   True),
      ("Feb 2026",  "New chair elected, Chile",     False)]
DOTX, TOP, STEP = 7.60, 1.98, 1.02
line_v(s, DOTX, TOP, TOP + STEP * 4, RULE, 1.5)
for i, (date, event, failed) in enumerate(TL):
    cy = TOP + STEP * i
    dot(s, DOTX, cy, 0.24 if failed else 0.14, AMBER_DEEP if failed else NAVY)
    # Both runs are middle-anchored in equal boxes centred on the dot, so the
    # 13pt date and the 16pt event share an optical centre with the dot.
    text(s, 7.78, cy - 0.20, 0.92, 0.40, date, size=LABEL, color=SLATE,
         bold=True, line=1.1, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 9.02, cy - 0.20, 3.71, 0.40, event, size=BODY,
         color=AMBER_DEEP if failed else NAVY, bold=failed, line=1.15,
         anchor=MSO_ANCHOR.MIDDLE)
source(s, "Sources: UNEP INC process; IISD Earth Negotiations Bulletin.")
notes(s, "In March 2022 the UN agreed to negotiate a binding treaty covering "
         "the full life cycle of plastic. Production, use, disposal. Five "
         "rounds of talks, finishing December 2024. That deadline came and "
         "went. So did the next one.")

# ================================================================ SLIDE 3
s = prs.slides.add_slide(BLANK)
bg(s, BONE)
text(s, M, 0.60, FULL, 0.72, "Both rounds died on the same split",
     size=TITLE, font=SERIF, bold=True, color=NAVY, line=1.1)
rect(s, 0, 1.65, 6.40, 5.85, fill=NAVY,
     descr="Navy panel: the High Ambition Coalition side of the split")

LP, LW = M, 5.20        # left panel text column
RP, RW = 7.00, 5.7333   # right panel text column
# Both section labels are muted so neither side outranks the other; 234 is
# the slide's single accent.
eyebrow(s, LP, 2.00, LW, "High ambition coalition", SLATE_LIGHT)
text(s, LP, 2.42, LW, 0.52, "70+ countries", size=HEAD, font=SERIF, bold=True,
     color=BONE, line=1.1)
text(s, LP, 3.08, LW, 0.95, [
        [("Panama, Rwanda, the EU. They want binding caps on ", {}),
         ("[production]", {"bold": True}), (".", {})]],
     size=BODY, color=BONE)
text(s, LP + 0.03, 5.10, LW, 1.10, "\u201cWe will not betray future generations.\u201d",
     size=HEAD, font=SERIF, italic=True, color=BONE, line=1.2)
source(s, "Panama's delegate, Geneva, Aug 2025. Membership: HAC, Aug 2025.",
       x=LP, y=6.55, w=LW, color=RULE)

eyebrow(s, RP, 2.00, RW, "Like-minded group", SLATE)
text(s, RP, 2.42, RW, 0.52, "Saudi Arabia, Russia", size=HEAD, font=SERIF,
     bold=True, color=NAVY, line=1.1)
text(s, RP, 3.08, RW, 0.95,
     "Oil-producing states. They want a treaty about bins and recycling. "
     "Waste only.", size=BODY, color=NAVY)
text(s, RP, 4.20, 3.40, 1.32, "234", size=BIGSTAT, font=SERIF, bold=True,
     color=AMBER_DEEP, line=1.0,
     descr="234 fossil-fuel and chemical lobbyists registered at the Geneva talks")
# Navy, not slate: this line is the stat's payload, not its citation.
text(s, RP, 5.55, RW, 0.85,
     "fossil-fuel and chemical lobbyists registered at Geneva. More people "
     "than the entire EU delegation.", size=BODY, color=NAVY)
source(s, "Source: CIEL analysis of the INC-5.2 participant list, Aug 2025",
       x=RP, y=6.55, w=RW)
notes(s, "Both rounds collapsed on the same split. On one side, around seventy "
         "countries wanting binding limits on how much plastic gets made. On "
         "the other, oil-producing states who wanted a treaty about bins and "
         "recycling only. Two hundred and thirty-four fossil-fuel lobbyists "
         "registered at Geneva. That was more people than the entire European "
         "Union sent. Panama's delegate said: we will not betray future "
         "generations. The talks ended with no deal. "
         "[PAUSE two seconds after the number.]")

# ================================================================ SLIDE 4
s = prs.slides.add_slide(BLANK)
bg(s, BONE)
text(s, M, 1.20, FULL, 1.55, [
        [("Is plastic a waste problem", {})],
        [("or a ", {}), ("[production]", {"color": AMBER_DEEP}),
         (" problem?", {})]],
     size=TITLE, font=SERIF, bold=True, color=NAVY, line=1.15,
     align=PP_ALIGN.CENTER)

# 0.50" gutter between the boxes, not the 0.20" grid gutter: on a slide whose
# whole argument is a binary, the two options must not read as one slab.
BW, BH, PAD = 4.79, 1.95, 0.40
BX, BX2, BY = col(2), 6.9167, 3.50
rect(s, BX, BY, BW, BH, line_color=SLATE, line_w=1.25,
     descr="Option box: plastic treated as a waste problem")
eyebrow(s, BX + PAD, BY + PAD, BW - 2 * PAD, "Waste", SLATE)
text(s, BX + PAD, BY + PAD + 0.42, BW - 2 * PAD, 1.05,
     "Clean it up. Recycle better. Blame the consumer.", size=BODY, color=NAVY)

# The endorsed option carries a 2.4x heavier stroke, so the pair is not a
# matched set of two equally weighted cards.
rect(s, BX2, BY, BW, BH, line_color=AMBER_DEEP, line_w=3.0,
     descr="Option box: plastic treated as a production problem")
eyebrow(s, BX2 + PAD, BY + PAD, BW - 2 * PAD, "Production", AMBER_DEEP)
text(s, BX2 + PAD, BY + PAD + 0.42, BW - 2 * PAD, 1.05,
     "Cap it at source. Ban toxic additives. Sue the maker.", size=BODY,
     color=NAVY)

source(s, "154 square brackets remain in the chair's August 2026 draft text "
          "(UNEP). Each one is a fight nobody has won yet.",
       x=BX, y=5.95, w=FULL - (BX - M))
notes(s, "Everything comes down to one question. Is plastic a waste problem or "
         "a production problem? If it is waste, you fix it with recycling and "
         "cleanups. If it is production, you have to tell oil companies to make "
         "less. That is why this is hard. And the science moved the argument. "
         "Researchers found microplastics in human artery plaque and in human "
         "brains. Plastic stopped being about turtles and started being about "
         "us. Those square brackets are the motif of this deck - every one is "
         "something countries still argue about.")

# ================================================================ SLIDE 5
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
text(s, M, 0.60, FULL, 0.72, "Losing twice radicalised the movement",
     size=TITLE, font=SERIF, bold=True, color=BONE, line=1.1)

# Evidence hangs off rows 2 and 4 only. The asymmetry is the point: four
# identically weighted rows is a table, and a table asserts without showing.
ROWS = [("Trusting UN consensus",     [("Demanding majority votes", {})], None),
        ("Voluntary company pledges", [("Binding treaty law", {})],
         "Coca-Cola cut its recycled-content target weeks after Busan"),
        ("Cleanup and recycling",     [("[production]", {"color": AMBER}),
                                       (" caps", {})], None),
        ("Lobbying governments",      [("Suing corporations", {})],
         "California v. ExxonMobil, filed September 2024")]
# Axis on the slide centre (6.667"): the two columns mirror each other.
# 26pt, not the 20pt body step: at 20pt the block filled 27% of the slide and
# floated with ~3in of dead navy either side, sharing an edge with nothing.
ROW5 = 26
AXL, AXR = 5.97, 7.37          # inner edges; 0.40" clear of the arrow each side
RTOP, RSTEP = 2.30, 1.09
eyebrow(s, M, 1.90, AXL - M, "Stopped", SLATE_LIGHT, align=PP_ALIGN.RIGHT)
eyebrow(s, AXR, 1.90, 12.7333 - AXR, "Started", BONE)
row_shapes = []
for i, (old, new, evidence) in enumerate(ROWS):
    ry = RTOP + RSTEP * i
    group = [text(s, M, ry, AXL - M, 0.55, old, size=ROW5, color=SLATE_LIGHT,
                  align=PP_ALIGN.RIGHT, line=1.15),
             # adj: slim shaft, proper head. The default is a 12px slab with a
             # 10px stub, which reads as the accent stripe the system bans.
             rect(s, 6.37, ry + 0.113, 0.60, 0.20, fill=SLATE_LIGHT,
                  shape=MSO_SHAPE.RIGHT_ARROW, adj=(0.25, 0.50),
                  descr="changed to"),
             text(s, AXR, ry, 12.7333 - AXR, 0.55,
                  [[(t, {**o, "bold": True}) for t, o in new]],
                  size=ROW5, color=BONE, line=1.15)]
    if evidence:
        group.append(text(s, AXR, ry + 0.48, 12.7333 - AXR, 0.26, evidence,
                          size=11, color=RULE, line=1.1))
    row_shapes.append(group)

if ANIMATE:
    add_click_build(s, row_shapes)
source(s, "Sources: CIEL; Coca-Cola 2030 target revision, Dec 2024; "
          "California Attorney General.", color=RULE, y=6.62)
notes(s, "So here is what changed. Four shifts. One, the movement stopped "
         "trusting consensus. When any one country can veto everything, "
         "campaigners now demand majority voting instead. Two, it stopped "
         "trusting company promises - weeks after Busan failed, Coca-Cola "
         "quietly weakened its recycled plastic targets. Three, it moved "
         "upstream, from cleaning beaches to capping production. And four, it "
         "went to court: California is suing ExxonMobil for telling the public "
         "that plastic was recyclable when it knew most of it was not. The "
         "movement did not get quieter after losing. It got harder to deal "
         "with. [Slow down. One sentence per row.]")

# ================================================================ SLIDE 6
s = prs.slides.add_slide(BLANK)
bg(s, BONE)
text(s, M, 0.60, FULL, 0.72, "Alive, but shrinking", size=TITLE, font=SERIF,
     bold=True, color=NAVY, line=1.1)
eyebrow(s, M, 2.05, 7.90, "What the chair took out", SLATE)
# Three runs: the strike crosses the word only, never the brackets.
text(s, M, 2.45, 7.90, 1.30, [[("[", {}),
                               ("production", {"strike": True}),
                               ("]", {})]],
     size=72, font=SERIF, bold=True, color=AMBER_DEEP, line=1.0,
     descr="The word production, struck through: cut from the August 2026 draft")
text(s, 7.20, 2.55, 4.40, 1.55,
     "Cut from the chair's August 2026 draft. France called it disappointing.",
     size=BODY, color=NAVY)

# The accent marks Aug 2026 - the moment the word was cut - not the furthest
# future date. Every dot is the same size so the labels stay flush left.
TL6 = [("Feb 2026", "New chair, Julio Cordano", False),
       ("Jul 2026", "Talks restart, Nairobi",   False),
       ("Aug 2026", "Draft drops production",   True),
       ("Sep 2026", "Bangkok session",          False),
       ("Mar 2027", "Formal talks resume",      False)]
T6Y, T6X0, T6P, T6W = 4.68, M, 2.4333, 2.2333
line_h(s, T6X0 + 0.075, T6X0 + 0.075 + T6P * 4, T6Y, RULE, 1.5)
for i, (date, event, key) in enumerate(TL6):
    cx = T6X0 + 0.075 + T6P * i
    dot(s, cx, T6Y, 0.15, AMBER_DEEP if key else NAVY)
    text(s, T6X0 + T6P * i, T6Y + 0.25, T6W, 0.26, date, size=11, color=SLATE,
         bold=True, line=1.1)
    text(s, T6X0 + T6P * i, T6Y + 0.55, T6W, 0.75, event, size=BODY,
         color=NAVY, bold=key, line=1.15)
source(s, "Sources: IISD Earth Negotiations Bulletin; UNEP; Climate Home News.",
       y=6.38)
notes(s, "February this year, countries elected a new chair, Julio Cordano of "
         "Chile. He restarted talks in Nairobi in July. But he kept production "
         "off the agenda, and the draft he published two weeks ago leaves "
         "production out entirely. France called it disappointing. Formal "
         "negotiations now happen in March 2027. So the process survived. The "
         "ambition is shrinking.")

# ================================================================ SLIDE 7
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
text(s, M, 3.02, FULL, 0.80, [
        [("Is a ", {}), ("[weak]", {"color": AMBER}),
         (" treaty better than no treaty?", {})]],
     size=TITLE, font=SERIF, bold=True, color=BONE, line=1.1,
     align=PP_ALIGN.CENTER)
# Only what the audience was actually shown. The microplastics research is
# spoken over slide 4, so it is cited in the notes, not here.
source(s, "Sources: UNEP; IISD Earth Negotiations Bulletin; CIEL; OECD Global "
          "Plastics Outlook 2022; Climate Home News.", color=RULE)
notes(s, "The movement's own answer is no. Its slogan now is that no treaty "
         "beats a bad treaty. But that is a gamble, because while they hold "
         "out, production keeps rising. So I will leave you with the question "
         "they are arguing over right now. Is a weak treaty better than no "
         "treaty at all?")

# ---------------------------------------------------------------- metadata
cp = prs.core_properties
cp.title = "The Treaty That Failed Twice"
cp.subject = "The Global Plastics Treaty and what failing twice did to the movement"
cp.author = "Cesco Cugliari"
cp.keywords = "plastics treaty, UNEP, INC, production cap"

prs.save(OUT)
print("wrote", OUT)
